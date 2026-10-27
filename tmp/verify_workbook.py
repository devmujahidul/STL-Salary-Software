from pathlib import Path

import openpyxl

for path in [
    Path(r"D:\MERND\output\Salary_Attendance_August_2026.xlsx"),
    Path(r"D:\MERND\output\Salary_Attendance_August_2026.xlsm"),
]:
    print("\n==", path.name, "==")
    wb = openpyxl.load_workbook(path, data_only=False)
    print("sheets:", [(ws.title, ws.sheet_state) for ws in wb.worksheets])
    for ws in wb.worksheets:
        print(
            f"{ws.title}: dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}"
        )
        print(
            f"  page: paper={ws.page_setup.paperSize} orient={ws.page_setup.orientation} "
            f"fitW={ws.page_setup.fitToWidth} fitH={ws.page_setup.fitToHeight} "
            f"scale={ws.page_setup.scale}"
        )
        print(
            f"  margins: L={ws.page_margins.left} R={ws.page_margins.right} "
            f"T={ws.page_margins.top} B={ws.page_margins.bottom}"
        )
        print(f"  print_area={ws.print_area} print_titles={ws.print_title_rows}")
        print(f"  freeze={ws.freeze_panes} autofilter={ws.auto_filter.ref}")
        print(f"  A4={ws['A4'].value!r}")
        print(f"  row6_height={ws.row_dimensions[6].height} row7_height={ws.row_dimensions[7].height}")
        print(f"  data_validations={len(ws.data_validations.dataValidation)}")
        for dv in ws.data_validations.dataValidation:
            print(f"    dv: {dv.type} {dv.formula1} cells={dv.sqref}")
        print(f"  hidden cols AU:AX = {[ws.column_dimensions[c].hidden for c in ('AU','AV','AW','AX')]}")
    print("defined names:", sorted(wb.defined_names.keys())[:12], "... total", len(wb.defined_names))
