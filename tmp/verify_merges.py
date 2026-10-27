"""Confirm the vertical merges in each worker block (SL/Name/Id/Designation,
the 9 monthly summary columns P..H and Signature span the 3 rows per worker)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

date_count = 31
identity = 4
# 1-indexed columns: dates occupy 5..identity+date_count
summary_start = identity + date_count + 1
expected_cols = {
    get_column_letter(c) for c in range(1, identity + 1)
} | {
    get_column_letter(c) for c in range(summary_start, summary_start + 9)
} | {
    get_column_letter(summary_start + 9)
}


def check(workbook_path: str) -> bool:
    wb = load_workbook(workbook_path)
    ok = True
    for name in wb.sheetnames:
        if name in {"Lists", "_Lists"}:
            continue
        ws = wb[name]
        merges = [m for m in ws.merged_cells.ranges if m.min_row == 6 and m.max_row == 8]
        cols = sorted(get_column_letter(m.min_col) for m in merges)
        missing = sorted(expected_cols - set(cols))
        print(f"{name}: merged block cols={cols}")
        if missing:
            print(f"  MISSING: {missing}")
            ok = False
    print("MERGE CHECK PASSED" if ok else "MERGE CHECK FAILED")
    return ok


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else str(
        Path(__file__).resolve().parent.parent / "output" / "Salary_Attendance_August_2026.xlsx"
    )
    sys.exit(0 if check(path) else 1)
