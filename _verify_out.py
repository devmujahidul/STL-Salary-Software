from openpyxl import load_workbook
from collections import Counter

path = r"D:\MERND\output\Attendance_August_2026.xlsx"
wb = load_workbook(path)
print("sheets", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n====", name)
    print("paper", ws.page_setup.paperSize, "orient", ws.page_setup.orientation)
    print("fit", ws.page_setup.fitToPage, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight)
    print("margins", ws.page_margins.left, ws.page_margins.right, ws.page_margins.top, ws.page_margins.bottom)
    print("print_title", ws.print_title_rows, ws.print_title_cols)
    print("print_area", ws.print_area)
    print("auto", ws.auto_filter.ref)
    print("freeze", ws.freeze_panes)
    print("A1-A4", [ws.cell(r, 1).value for r in range(1, 5)])
    if name == "_Lists":
        continue
    print("header", [ws.cell(5, c).value for c in range(1, ws.max_column + 1)])
    ids = []
    subs = Counter()
    for r in range(6, ws.max_row + 1):
        emp = ws.cell(r, 3).value
        sub = ws.cell(r, ws.max_column).value
        if emp:
            ids.append(emp)
            subs[sub] += 1
    print("workers", len(ids), "prefixes", Counter(x.split("-")[0] for x in ids))
    print("subsections", dict(subs))
    print("row6", [ws.cell(6, c).value for c in range(1, 10)], "sub", ws.cell(6, ws.max_column).value)

ws = wb["FIXED"]
print("\nA4 formula", ws["A4"].value)
print("dv count", len(ws.data_validations.dataValidation))
cl = [ws.cell(r, 3).value for r in range(6, ws.max_row + 1) if str(ws.cell(r, 3).value).startswith("CL-")]
print("CL ids", len(cl), cl[:5])
cl_sub = {
    ws.cell(r, ws.max_column).value
    for r in range(6, ws.max_row + 1)
    if str(ws.cell(r, 3).value).startswith("CL-")
}
print("CL subsections", cl_sub)
